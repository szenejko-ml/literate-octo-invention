"""Aggregate ONLY held-out predictions; separate operating errors and phenotypes.

No fitting of predictors and no threshold optimization occurs in this module.
Phenotype tests use disjoint A-only/B-only error groups; comparator denominators
and anomaly/CHD are never in the primary phenotype test family.
"""
from __future__ import annotations
from itertools import combinations
from pathlib import Path
import argparse
import json

import numpy as np
import pandas as pd
from scipy import stats, optimize
from scipy.special import expit, logit
from sklearn.metrics import roc_auc_score, average_precision_score, brier_score_loss
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from doctoral_models import Protocol, BINARY, CONTROLS, feature_names, override_mask


def bh(p):
    p=np.asarray(p,float); out=np.full(len(p),np.nan); keep=np.flatnonzero(np.isfinite(p))
    if len(keep):
        order=np.argsort(p[keep]); ranked=p[keep][order]; m=len(ranked)
        q=np.minimum.accumulate((ranked*m/np.arange(1,m+1))[::-1])[::-1].clip(0,1)
        mapped=np.empty(m); mapped[order]=q; out[keep]=mapped
    return out


def exact_ci(k,n):
    if not n: return np.nan,np.nan
    return (0.0 if not k else float(stats.beta.ppf(.025,k,n-k+1)),
            1.0 if k==n else float(stats.beta.ppf(.975,k+1,n-k)))


def cal_summary(y,p):
    # Intercept-only calibration-in-the-large and joint calibration slope/intercept
    # are deliberately not conflated. Descriptive estimates on held-out predictions.
    y=np.asarray(y,int); p=np.clip(np.asarray(p,float),1e-6,1-1e-6)
    if len(np.unique(y))<2 or len(np.unique(p))<2: return {}
    x=logit(p); design=np.column_stack([np.ones(len(y)),x])
    def fun(beta):
        eta=design@beta
        return float(np.sum(np.logaddexp(0,eta)-y*eta))
    def jac(beta): return design.T@(expit(design@beta)-y)
    opt=optimize.minimize(fun,[0,1],jac=jac,method='BFGS',options={'maxiter':1000,'gtol':1e-6})
    if not np.isfinite(opt.x).all() or np.max(np.abs(jac(opt.x)))>1e-3:
        return {'calibration_status':'unstable_or_failed'}
    fitted=expit(design@opt.x); w=fitted*(1-fitted)
    try: se=np.sqrt(np.diag(np.linalg.inv((design.T*w)@design)))
    except np.linalg.LinAlgError: se=np.array([np.nan,np.nan])
    citl=optimize.brentq(lambda a:np.sum(expit(x+a)-y),-60,60)
    out={'calibration_status':'estimated','calibration_in_the_large':citl}
    for i,key in enumerate(['joint_calibration_intercept','calibration_slope']):
        out[key]=float(opt.x[i]); out[key+'_CI_low']=float(opt.x[i]-1.96*se[i]); out[key+'_CI_high']=float(opt.x[i]+1.96*se[i])
    return out


def auc_influence(y,s):
    y=np.asarray(y,int); s=np.asarray(s,float)
    a=s[y==1]; b=s[y==0]; m=len(a); n=len(b)
    if min(m,n)<2: return np.nan,None,None
    r=stats.rankdata(s,method='average')
    vp=(r[y==1]-stats.rankdata(a,method='average'))/n
    vn=1-(r[y==0]-stats.rankdata(b,method='average'))/m
    return float(vp.mean()),vp,vn


def auc_compare(y,a,b):
    aa,ap,an=auc_influence(y,a); ba,bp,bn=auc_influence(y,b)
    if ap is None: return {}
    diff=ba-aa
    var=np.var(bp-ap,ddof=1)/len(ap)+np.var(bn-an,ddof=1)/len(an)
    se=float(np.sqrt(max(var,0)))
    pvalue=1.0 if se==0 and diff==0 else (0.0 if se==0 else float(2*stats.norm.sf(abs(diff/se))))
    return {'AUC_A':aa,'AUC_B':ba,'AUC_difference_B_minus_A':diff,
            'AUC_difference_CI_low':diff-1.96*se,'AUC_difference_CI_high':diff+1.96*se,
            'AUC_DeLong_p':pvalue}


def rates(y,z):
    y=np.asarray(y,int); z=np.asarray(z,object)
    eu=y==1; non=y==0; low=z=='low'; high=z=='high'; mid=z=='intermediate'
    return {'n':len(y),'n_euploid':int(eu.sum()),'n_non_euploid':int(non.sum()),
            'low_n':int(low.sum()),'intermediate_n':int(mid.sum()),'high_n':int(high.sum()),
            'low_euploid':int((low&eu).sum()),'low_non_euploid':int((low&non).sum()),
            'intermediate_euploid':int((mid&eu).sum()),'intermediate_non_euploid':int((mid&non).sum()),
            'high_euploid':int((high&eu).sum()),'high_non_euploid':int((high&non).sum()),
            'euploid_nonlow_n':int((~low&eu).sum())}


def performance(data,method,p,variant='policy'):
    zcol=('zone_' if variant=='policy' else 'zone_model_only_')+method
    good=data[zcol].isin(['low','intermediate','high'])
    d=data.loc[good]; y=d.target.to_numpy(int); z=d[zcol].to_numpy()
    out={'method':method,'variant':variant,'n_source_cohort':len(data),'n_unavailable':int((~good).sum()),**rates(y,z)}
    for key,num,den in [
        ('low_coverage','low_n','n'),('false_reassurance_among_non_euploid','low_non_euploid','n_non_euploid'),
        ('residual_non_euploid_among_low','low_non_euploid','low_n'),
        ('euploid_predictive_value_of_low','low_euploid','low_n'),
        ('high_sensitivity','high_non_euploid','n_non_euploid'),('high_fpr','high_euploid','n_euploid'),
        ('high_ppv','high_non_euploid','high_n'),('euploid_nonlow_rate','euploid_nonlow_n','n_euploid')]:
        k=out[num]; n=out[den]; out[key]=k/n if n else np.nan
        out[key+'_CI_low'],out[key+'_CI_high']=exact_ci(k,n)
    if 'prob_'+method in d:
        prob=d['prob_'+method].to_numpy(float)
        auc,vp,vn=auc_influence(y,prob)
        out['AUC']=auc
        if vp is not None:
            se=np.sqrt(np.var(vp,ddof=1)/len(vp)+np.var(vn,ddof=1)/len(vn))
            out['AUC_CI_low']=max(0,auc-1.96*se); out['AUC_CI_high']=min(1,auc+1.96*se)
        out['average_precision_non_euploid']=average_precision_score(1-y,1-prob)
        out['brier_euploid']=brier_score_loss(y,prob)
        out.update(cal_summary(y,prob))
    else:
        score=d['score_'+method].to_numpy(float)
        out['AUC']=roc_auc_score(y,score)
        out['calibration_status']='not_applicable_comparator_rank_score_is_not_joint_probability'
    out['meets_low_error_point_estimate']=out['false_reassurance_among_non_euploid']<=p.low_error_cap+1e-12
    out['meets_high_fpr_point_estimate']=out['high_fpr']<=p.high_fpr_cap+1e-12
    out['nondegenerate_low_zone']=out['low_coverage']>=p.minimum_low_coverage
    out['meets_all_operating_point_criteria']=all(out[k] for k in ['meets_low_error_point_estimate','meets_high_fpr_point_estimate','nondegenerate_low_zone'])
    return out


def error_pairs(data,methods,p):
    rows=[]; comparisons=[]
    for a,b in combinations(methods,2):
        available=data['zone_'+a].isin(['low','intermediate','high'])&data['zone_'+b].isin(['low','intermediate','high'])
        d=data.loc[available]; y=d.target.to_numpy(int)
        za=d['zone_'+a].to_numpy(); zb=d['zone_'+b].to_numpy()
        for outcome in ['euploid_nonlow','non_euploid_low']:
            population=y==(1 if outcome=='euploid_nonlow' else 0)
            ea=za!='low' if outcome=='euploid_nonlow' else za=='low'
            eb=zb!='low' if outcome=='euploid_nonlow' else zb=='low'
            aa=ea[population]; bb=eb[population]
            both=int((aa&bb).sum()); aonly=int((aa&~bb).sum()); bonly=int((~aa&bb).sum())
            neither=int((~aa&~bb).sum()); n=len(aa); nd=aonly+bonly; union=both+nd
            row={'outcome':outcome,'A':a,'B':b,'n_class':n,'A_errors':both+aonly,'B_errors':both+bonly,
                 'both':both,'A_only':aonly,'B_only':bonly,'neither':neither,'union':union,
                 'Jaccard':both/union if union else np.nan,'overlap_as_fraction_of_A':both/(both+aonly) if both+aonly else np.nan,
                 'overlap_as_fraction_of_B':both/(both+bonly) if both+bonly else np.nan,
                 'difference_B_minus_A':(bonly-aonly)/n if n else np.nan,
                 'McNemar_p':float(stats.binomtest(aonly,nd,.5).pvalue) if nd else 1.0}
            if n and p.bootstrap_reps:
                rng=np.random.default_rng(p.seed)
                samp=rng.multinomial(n,np.array([aonly,bonly,n-aonly-bonly])/n,size=p.bootstrap_reps)
                diffs=(samp[:,1]-samp[:,0])/n
                row['difference_CI_low'],row['difference_CI_high']=np.quantile(diffs,[.025,.975])
            rows.append(row)
            # Disjoint error sets only. Never compare overlapping sets as independent samples.
            comparisons.append((outcome,a,b,d.loc[population&ea&~eb],d.loc[population&~ea&eb]))
    table=pd.DataFrame(rows)
    if len(table):
        table['McNemar_q']=table.groupby('outcome')['McNemar_p'].transform(lambda s:bh(s))
    return table,comparisons


def group_summary(group,feature):
    x=pd.to_numeric(group[feature],errors='coerce').dropna().to_numpy(float)
    if not len(x): return {'n':0}
    if feature in BINARY:
        return {'n':len(x),'events':int((x==1).sum()),'proportion':float((x==1).mean())}
    if feature=='DV':
        return {'n':len(x),**{f'n_{k}':int((x==k).sum()) for k in [0,1,2]},
                **{f'proportion_{k}':float((x==k).mean()) for k in [0,1,2]}}
    q=np.quantile(x,[.25,.5,.75]); return {'n':len(x),'q1':q[0],'median':q[1],'q3':q[2]}


def fisher_freeman_halton(table, n_resamples=9999, seed=42):
    """Two-sided conditional 2x3 test; exact for small margins, MC otherwise.

    Uses probability ordering, as Fisher's two-sided test does. Drawing tables
    directly avoids a reshape failure in some SciPy MonteCarloMethod batches.
    """
    tab=np.asarray(table,dtype=int)
    if tab.ndim!=2 or tab.shape[0]!=2 or (tab<0).any():
        raise ValueError('Expected a non-negative 2-by-k count table')
    tab=tab[:,tab.sum(axis=0)>0]
    if tab.shape[1]<2:
        return 1.0,'constant_category'
    if tab.shape[1]==2:
        return float(stats.fisher_exact(tab).pvalue),'two_sided_Fisher_exact_2x2'
    if tab.shape[1]!=3:
        raise ValueError('DV has three categories; no higher-dimensional table is expected')
    if tab[0].sum()>tab[1].sum():
        tab=tab[::-1]
    row=tab.sum(axis=1); col=tab.sum(axis=0)
    distribution=stats.random_table(row,col)
    observed=float(distribution.logpmf(tab))
    if row[0]<=20:
        tables=[]
        for x in range(min(int(row[0]),int(col[0]))+1):
            for z in range(min(int(row[0])-x,int(col[1]))+1):
                third=int(row[0])-x-z
                if 0<=third<=col[2]:
                    top=np.array([x,z,third])
                    tables.append(np.stack([top,col-top]))
        logprob=np.asarray(distribution.logpmf(np.asarray(tables)))
        value=float(np.exp(logprob[logprob<=observed+1e-12]).sum())
        return min(1.,value),'Fisher_Freeman_Halton_exact'
    rng=np.random.default_rng(seed); extreme=0
    for start in range(0,n_resamples,256):
        size=min(256,n_resamples-start)
        draws=distribution.rvs(size=size,random_state=rng)
        extreme+=int(np.sum(distribution.logpmf(draws)<=observed+1e-12))
    return (extreme+1)/(n_resamples+1),'Fisher_Freeman_Halton_MonteCarlo'


def phenotype_tests(comparisons,p):
    rows=[]; features=feature_names(p.nt_representation)
    for outcome,a,b,ga,gb in comparisons:
        for feature in features:
            xa=pd.to_numeric(ga[feature],errors='coerce').dropna().to_numpy(float)
            xb=pd.to_numeric(gb[feature],errors='coerce').dropna().to_numpy(float)
            row={'outcome':outcome,'A':a,'B':b,'feature':feature,'n_A_only':len(xa),'n_B_only':len(xb),
                 'comparison':'Disjoint A-only versus B-only errors; conditional exploratory phenotype analysis',
                 **{'A_'+k:v for k,v in group_summary(ga,feature).items()},
                 **{'B_'+k:v for k,v in group_summary(gb,feature).items()},'p':np.nan}
            if min(len(xa),len(xb))<2:
                row['test']='not_tested_fewer_than_2_cases_in_one_group'
            elif feature in BINARY:
                k1=int((xa==1).sum()); k2=int((xb==1).sum())
                row['p']=float(stats.fisher_exact([[k1,len(xa)-k1],[k2,len(xb)-k2]])[1])
                row['effect_B_minus_A']=k2/len(xb)-k1/len(xa)
                row['effect_type']='risk_difference'; row['test']='two_sided_Fisher_exact'
            elif feature=='DV':
                tab=np.array([[(xa==k).sum() for k in [0,1,2]],[(xb==k).sum() for k in [0,1,2]]])
                tab=tab[:,tab.sum(axis=0)>0]
                if tab.shape[1]<2:
                    row['p']=1.0; row['test']='constant_category'
                else:
                    chi,prob,_,expected=stats.chi2_contingency(tab,correction=False)
                    if (expected<5).any():
                        row['p'],row['test']=fisher_freeman_halton(tab,p.phenotype_permutations,p.seed)
                        if row['test'].endswith('MonteCarlo'):
                            row['MonteCarlo_resamples']=p.phenotype_permutations
                    else:
                        row['p']=float(prob); row['test']='Pearson_chi_square_2x3'
                    row['effect_B_minus_A']=float(np.sqrt(chi/tab.sum()))
                    row['effect_type']='Cramers_V_unsigned'
            else:
                if min(len(xa),len(xb))<8:
                    method=stats.PermutationMethod(n_resamples=p.phenotype_permutations,batch=128,rng=np.random.default_rng(p.seed))
                    result=stats.mannwhitneyu(xa,xb,alternative='two-sided',method=method)
                    row['test']='Mann_Whitney_permutation_small_sample'
                else:
                    result=stats.mannwhitneyu(xa,xb,alternative='two-sided',method='asymptotic')
                    row['test']='Mann_Whitney_tie_corrected'
                row['p']=float(result.pvalue)
                row['effect_B_minus_A']=float(1-2*result.statistic/(len(xa)*len(xb)))
                row['effect_type']='rank_biserial_B_minus_A'
            row['small_sample_caution']=min(len(xa),len(xb))<10
            rows.append(row)
    out=pd.DataFrame(rows)
    if len(out):
        out['q_BH']=out.groupby('outcome')['p'].transform(lambda s:bh(s))
        out['nominal_p_lt_005']=out['p']<.05; out['FDR_q_lt_005']=out['q_BH']<.05
    return out


def profiles(data,methods,p):
    rows=[]
    for method in methods:
        z=data['zone_'+method]
        for outcome,mask in [('euploid_nonlow',data.target.eq(1)&z.isin(['intermediate','high'])),
                             ('non_euploid_low',data.target.eq(0)&z.eq('low'))]:
            for stratum in ['all','without_override','with_override']:
                force=override_mask(data)
                m=mask.to_numpy()&(np.ones(len(data),bool) if stratum=='all' else (force if stratum=='with_override' else ~force))
                group=data.loc[m]
                for f in feature_names(p.nt_representation)+(['NT_raw'] if p.nt_representation!='raw' else []):
                    rows.append({'method':method,'outcome':outcome,'override_stratum':stratum,
                                 'group_n':len(group),'feature':f,**group_summary(group,f)})
    return pd.DataFrame(rows)


def export_workbook(tables,path):
    wb=Workbook(); wb.remove(wb.active)
    for title,frame in tables.items():
        ws=wb.create_sheet(title[:31]); ws.sheet_view.showGridLines=False
        columns=list(frame.columns); ws.append(columns)
        for row in frame.itertuples(index=False,name=None):
            ws.append([None if (v is None or (isinstance(v,(float,np.floating)) and not np.isfinite(v)))
                       else (v.item() if isinstance(v,np.generic) else v) for v in row])
        ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor='20384D'); cell.font=Font(color='FFFFFF',bold=True)
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[1].height=58
        for j,name in enumerate(columns,1):
            ws.column_dimensions[get_column_letter(j)].width=min(38,max(14,len(str(name))*.72))
            for row in ws.iter_rows(min_row=2,min_col=j,max_col=j):
                cell=row[0]; cell.font=Font(name='Calibri',size=10,color='008000')
                cell.alignment=Alignment(vertical='center',wrap_text=True)
                if isinstance(cell.value,float):
                    cell.number_format='0.0000' if ('p'==name or name.endswith('_p') or 'q_BH'==name or name.endswith('_q')) else '0.000'
        ws.sheet_properties.pageSetUpPr.fitToPage=True
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    # Tables are imports of statistical-computation outputs, not Excel calculations.
    # Simple ratios remain fully reconstructable from accompanying integer counts.
    temp=path.with_name(path.stem+'.tmp.xlsx'); wb.save(temp); temp.replace(path)


def summarize_cohort(directory,p):
    directory=Path(directory); d=pd.read_csv(directory/'private'/'oof_predictions.csv.gz')
    methods=[c[5:] for c in d if c.startswith('zone_') and not c.startswith('zone_model_only_')]
    perf=pd.DataFrame([performance(d,m,p,v) for m in methods for v in ['policy','model_only']])
    # Paired inference uses the full policy. Override-only contribution is separate.
    pairs,comparisons=error_pairs(d,methods,p)
    pheno=phenotype_tests(comparisons,p)
    prof=profiles(d,methods,p)
    auc=[]
    for a,b in combinations(methods,2):
        ca=('prob_' if 'prob_'+a in d else 'score_')+a
        cb=('prob_' if 'prob_'+b in d else 'score_')+b
        valid=np.isfinite(d[ca])&np.isfinite(d[cb])
        yy=d.loc[valid,'target'].to_numpy(int)
        auc.append({'A':a,'B':b,'n_common':int(valid.sum()),**auc_compare(yy,d.loc[valid,ca],d.loc[valid,cb])})
    auc=pd.DataFrame(auc)
    if len(auc) and 'AUC_DeLong_p' in auc: auc['AUC_DeLong_q']=bh(auc['AUC_DeLong_p'])
    overrides=[]
    for m in methods:
        a=d['zone_model_only_'+m]; b=d['zone_'+m]
        for label in [0,1]:
            mask=d.target.eq(label)&a.notna()&b.notna()
            overrides.append({'method':m,'target':label,'new_high_due_to_override':int((mask&a.ne('high')&b.eq('high')).sum()),
                              'low_to_high_due_to_override':int((mask&a.eq('low')&b.eq('high')).sum()),
                              'interpretation':'Fixed thresholds, mechanical override contribution; not reoptimized no-override policy.'})
    # Intersections are descriptive and not a proposed ensemble policy.
    nonlinear=[m for m in ['random_forest','extra_trees','xgboost','catboost','lightgbm','tabpfn'] if m in methods]
    consensus=[]
    for label,event in [(1,'euploid_nonlow'),(0,'non_euploid_low')]:
        g=d.loc[d.target.eq(label)].copy()
        if nonlinear:
            votes=sum((g['zone_'+m].ne('low') if label==1 else g['zone_'+m].eq('low')).astype(int) for m in nonlinear)
            for k in range(len(nonlinear)+1):
                gg=g.loc[votes.eq(k)]
                for f in feature_names(p.nt_representation):
                    consensus.append({'outcome':event,'panel':','.join(nonlinear),'number_of_models':len(nonlinear),
                                      'number_agreeing':k,'group_n':len(gg),'feature':f,**group_summary(gg,f)})
    consensus=pd.DataFrame(consensus)
    notes=pd.DataFrame({'note':[
        'Performance is computed from outer-test predictions only. Full-cohort final fits are not evaluated here.',
        'Comparators use observed denominators only. Pairwise analyses restrict to records available for both methods.',
        'Euploid non-low = reassurance withheld; not automatically unnecessary clinical referral, especially with an anomaly.',
        'False reassurance = non-euploid in low risk. Its denominator is all eligible non-euploid outcomes.',
        'McNemar tests equality of marginal error probabilities, not patient identity or biological equivalence.',
        'Phenotype tests compare mutually exclusive A-only and B-only error groups. All are post-selection exploratory tests.',
        'Primary phenotype family: the 10 model input variables. Comparator denominators, controls and duplicate NT representations are excluded.',
        'BH families: each error outcome across all method pairs and all 10 phenotype variables. Nominal p and q are both retained.',
        'Calibration slopes/intercepts and their Wald intervals are descriptive pooled-OOF summaries, not external validation.',
        'No difference/equivalence can be inferred merely from q >= .05. Small false-reassurance groups can be uninformative.',
        'CST/CST+ versus ML differences combine predictors, outcomes, calibration and threshold policies; not isolated algorithm effects.',
        'Consensus counts use the available nonlinear panel, explicitly listed. They are NOT a best-model selection rule or validated ensemble.',
        'Patient-level bootstrap intervals condition on the fitted cross-validation predictions and do not refit the complete pipeline.',
        'Standardized or rank-based effect sizes accompany p values; significance is not a causal importance measure.'
    ]})
    tables={'readme':notes,'performance':perf,'paired_errors':pairs,'phenotype_tests':pheno,
            'error_profiles':prof,'AUC_pairs':auc,'override_contribution':pd.DataFrame(overrides),
            'descriptive_consensus':consensus}
    out=directory/'aggregate'; out.mkdir(parents=True,exist_ok=True)
    for name,table in tables.items(): table.to_csv(out/(name+'.csv'),index=False)
    export_workbook(tables,out/'doctoral_results.xlsx')


def sensitivity_comparison(root,p):
    root=Path(root); a=root/'all_years_shared'/'private'/'oof_predictions.csv.gz'
    b=root/'all_years_predictor_complete'/'private'/'oof_predictions.csv.gz'
    if not a.exists() or not b.exists(): return
    x=pd.read_csv(a); y=pd.read_csv(b)
    merged=x.merge(y,on='record_id',suffixes=('_shared','_broad'),validate='one_to_one')
    assert len(merged)==len(x)
    assert np.array_equal(merged.target_shared,merged.target_broad)
    assert np.array_equal(merged.outer_fold_shared,merged.outer_fold_broad)
    methods=[c[5:] for c in x if c.startswith('prob_') and c in y]
    rows=[]
    for m in methods:
        z1=merged['zone_'+m+'_shared']; z2=merged['zone_'+m+'_broad']
        for label,event in [(1,'euploid_nonlow'),(0,'non_euploid_low')]:
            mask=merged.target_shared.eq(label)
            ea=(z1.ne('low') if label==1 else z1.eq('low'))[mask]
            eb=(z2.ne('low') if label==1 else z2.eq('low'))[mask]
            ao=int((ea&~eb).sum()); bo=int((~ea&eb).sum()); both=int((ea&eb).sum())
            rows.append({'model':m,'outcome':event,'n_common':int(mask.sum()),'both':both,
                'shared_training_only':ao,'broader_training_only':bo,'McNemar_p':stats.binomtest(ao,ao+bo,.5).pvalue if ao+bo else 1.0,
                'interpretation':'Both policies evaluated on identical common records and identical outer folds; only eligible training cohort differs.'})
    table=pd.DataFrame(rows)
    if len(table):
        table['McNemar_q']=table.groupby('outcome')['McNemar_p'].transform(lambda s:bh(s))
        table.to_csv(root/'selection_sensitivity_common_records.csv',index=False)

if __name__=='__main__':
    a=argparse.ArgumentParser(); a.add_argument('--root',required=True)
    args=a.parse_args(); root=Path(args.root)
    p=Protocol(**json.loads((root/'protocol.json').read_text())['protocol'])
    for directory in root.iterdir():
        if (directory/'private'/'oof_predictions.csv.gz').exists(): summarize_cohort(directory,p)
    sensitivity_comparison(root,p)